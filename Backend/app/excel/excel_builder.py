"""
Excel export for extraction results.

- "Results" sheet: one row per extraction_result. Columns: Source File,
  then each non-table blueprint_field in sort_order (header = display_label),
  then Confidence Score, then Status.
- One additional sheet per table-type blueprint field (e.g. "Line Items",
  "Grades"): one row per item across all results, columns = Source File +
  the field's table_columns (so a 63-row marksheet lands as 63 real Excel
  rows, not one flattened cell).
"""
import logging
import re
import uuid
from pathlib import Path
from typing import Any, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select

from app.config import settings
from app.database import AsyncSessionLocal, BlueprintFieldORM, ExtractionResultORM, JobORM

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

_INVALID_SHEET_CHARS = re.compile(r"[:\\/?*\[\]]")


def _humanize(key: str) -> str:
    return key.replace("_", " ").strip().title()


def _safe_sheet_name(label: str, existing: set) -> str:
    """Excel sheet names: <=31 chars, no : \\ / ? * [ ], unique within the workbook."""
    base = _INVALID_SHEET_CHARS.sub(" ", label).strip() or "Table"
    base = base[:31]
    name = base
    n = 2
    while name.lower() in existing:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    existing.add(name.lower())
    return name


def _write_header(ws, headers: List[str]):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = 22
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def build_workbook(fields: List[BlueprintFieldORM], results: List[ExtractionResultORM]) -> Workbook:
    scalar_fields = [f for f in fields if f.field_type != "table"]
    table_fields = [f for f in fields if f.field_type == "table"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"

    headers = ["Source File"] + [f.display_label for f in scalar_fields] + ["Confidence Score", "Status"]
    _write_header(ws, headers)

    for r in results:
        data = r.data or {}
        row = [r.source_filename]
        row += [data.get(f.field_key) for f in scalar_fields]
        row += [r.confidence_score, r.status]
        ws.append(row)

    sheet_names: set = {"results"}
    for field in table_fields:
        columns = field.table_columns or []
        col_keys = [c.get("key") or c.get("field_key") for c in columns if c.get("key") or c.get("field_key")]

        table_ws = wb.create_sheet(_safe_sheet_name(field.display_label, sheet_names))
        col_headers = ["Source File"] + [_humanize(k) for k in col_keys]
        _write_header(table_ws, col_headers)

        for r in results:
            rows_value = (r.data or {}).get(field.field_key)
            if not isinstance(rows_value, list):
                continue
            for item in rows_value:
                if isinstance(item, dict):
                    if not col_keys:
                        # No table_columns recorded on the field — fall back to
                        # whatever keys actually showed up in the data.
                        col_keys = list(item.keys())
                        table_ws.delete_rows(1, table_ws.max_row)
                        _write_header(table_ws, ["Source File"] + [_humanize(k) for k in col_keys])
                    table_ws.append([r.source_filename] + [item.get(k) for k in col_keys])
                else:
                    table_ws.append([r.source_filename, item])

    return wb


async def build_excel(job_id: str) -> Path:
    output_dir = Path(settings.OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{job_id}.xlsx"

    async with AsyncSessionLocal() as session:
        job = await session.get(JobORM, uuid.UUID(job_id))
        fields: List[BlueprintFieldORM] = []
        if job and job.blueprint_id:
            field_result = await session.execute(
                select(BlueprintFieldORM)
                .where(BlueprintFieldORM.blueprint_id == job.blueprint_id)
                .order_by(BlueprintFieldORM.sort_order)
            )
            fields = list(field_result.scalars().all())

        result_rows = await session.execute(
            select(ExtractionResultORM)
            .where(ExtractionResultORM.job_id == uuid.UUID(job_id))
            .order_by(ExtractionResultORM.extracted_at)
        )
        results = list(result_rows.scalars().all())

    wb = build_workbook(fields, results)
    wb.save(str(output_path))
    logger.info(f"Excel saved: {output_path} ({len(results)} result(s))")
    return output_path
